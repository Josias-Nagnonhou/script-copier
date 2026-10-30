"""Kopi — récupère la transcription de plusieurs vidéos TikTok / YouTube
et les rassemble dans un document Word (Script 1, Script 2, ...)."""
import os, re, json, time, uuid, glob, shutil, threading, tempfile, subprocess, datetime
from concurrent.futures import ThreadPoolExecutor

import requests
import yt_dlp
from flask import Flask, request, jsonify, send_file, render_template, session, redirect
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or uuid.uuid4().hex
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 ** 3  # 2 Go par fichier

def find_groq_key():
    """Trouve la clé Groq même si le nom de la variable est mal écrit
    (espaces, minuscules, guillemets) ou si la clé a été collée ailleurs."""
    def clean(v):
        return (v or "").strip().strip('"').strip("'").strip()
    for k, v in os.environ.items():
        if "GROQ" in k.strip().upper() and clean(v):
            return clean(v)
    for k, v in os.environ.items():
        if clean(v).startswith("gsk_"):
            return clean(v)
        if clean(k).startswith("gsk_"):  # clé collée dans la case du nom
            return clean(k)
    return ""


GROQ_KEY = find_groq_key()
print("Clé Groq :", "trouvée" if GROQ_KEY else "ABSENTE", "| variables :",
      sorted(k for k in os.environ if k.isupper() and not k.startswith(("RENDER", "PYTHON", "PATH", "HOME", "LANG", "GPG", "HOSTNAME"))), flush=True)
GROQ_MODEL = os.environ.get("GROQ_MODEL", "whisper-large-v3-turbo")
APP_PASSWORD = os.environ.get("APP_PASSWORD", "").strip()
WORKERS = int(os.environ.get("WORKERS", "1"))  # serveur gratuit : 512 Mo, une vidéo à la fois
BASE = os.path.join(tempfile.gettempdir(), "script-copier")
os.makedirs(BASE, exist_ok=True)

COOKIES = {}
# état des cookies YouTube : absent / ok / expired (+ d'où ils viennent)
YT = {"status": "absent", "source": "", "updated": None}


def valid_yt_cookies(text):
    lines = [l for l in text.splitlines() if l.strip() and not l.startswith("#")]
    return any("youtube.com" in l and len(l.split("\t")) >= 7 for l in lines)


def install_yt_cookies(text, source):
    p = os.path.join(BASE, "yt_cookies.txt")  # copie modifiable (yt-dlp la met à jour)
    with open(p, "w") as f:
        f.write(text if text.startswith("# Netscape") else "# Netscape HTTP Cookie File\n" + text)
    COOKIES["YT_COOKIES"] = p
    YT.update(status="ok", source=source, updated=time.time())


def load_cookies():
    # 1) fichier secret Render (permanent) 2) variable d'environnement
    for p in sorted(glob.glob("/etc/secrets/*")):
        try:
            t = open(p, encoding="utf-8", errors="ignore").read()
        except OSError:
            continue
        if valid_yt_cookies(t):
            install_yt_cookies(t, "render")
            return
    for k, v in os.environ.items():
        if "COOKIE" in k.upper() and "TT" not in k.upper() and valid_yt_cookies(v):
            install_yt_cookies(v, "render")
            return


load_cookies()
if os.environ.get("TT_COOKIES"):
    _p = os.path.join(BASE, "tt_cookies.txt")
    with open(_p, "w") as f:
        f.write(os.environ["TT_COOKIES"])
    COOKIES["TT_COOKIES"] = _p
print("Cookies YouTube :", YT["status"], flush=True)

JOBS = {}
LOCK = threading.Lock()


# ---------------------------------------------------------------- helpers
def ydl_opts(platform, **extra):
    o = {"quiet": True, "no_warnings": True, "noprogress": True,
         "socket_timeout": 30, "retries": 3}
    ck = COOKIES.get("YT_COOKIES" if platform == "youtube" else "TT_COOKIES")
    if ck:
        o["cookiefile"] = ck
    o.update(extra)
    return o


def normalize_source(platform, raw, yt_tab):
    raw = raw.strip()
    if platform == "tiktok":
        if raw.startswith("http"):
            return raw
        return "https://www.tiktok.com/@" + raw.lstrip("@").strip()
    # youtube
    tab = "shorts" if yt_tab == "shorts" else "videos"
    if not raw.startswith("http"):
        return f"https://www.youtube.com/@{raw.lstrip('@').strip()}/{tab}"
    m = re.match(r"^(https?://(?:www\.|m\.)?youtube\.com/(?:@[^/?#]+|channel/[^/?#]+|c/[^/?#]+|user/[^/?#]+))/?(?:featured)?/?$", raw)
    if m:
        return f"{m.group(1)}/{tab}"
    return raw


def flatten(info):
    out = []
    for e in info.get("entries") or []:
        if not e:
            continue
        if e.get("entries"):
            out += flatten(e)
        else:
            out.append(e)
    return out


def entry_url(e, platform):
    u = e.get("webpage_url") or e.get("url") or ""
    if u.startswith("http"):
        return u
    vid = e.get("id") or u
    if platform == "youtube":
        return "https://www.youtube.com/watch?v=" + vid
    return u


def list_videos(platform, url, limit, sort):
    opts = ydl_opts(platform, extract_flat="in_playlist", skip_download=True, ignoreerrors=True)
    if sort == "recent" and limit:
        opts["playlistend"] = limit
    with yt_dlp.YoutubeDL(opts) as y:
        info = y.extract_info(url, download=False)
    if not info:  # relancer sans ignoreerrors pour obtenir la vraie raison
        opts.pop("ignoreerrors")
        try:
            with yt_dlp.YoutubeDL(opts) as y:
                info = y.extract_info(url, download=False)
        except Exception as e:
            raise RuntimeError(f"Compte ou lien introuvable ({e})")
    if not info:
        raise RuntimeError("Compte ou lien introuvable.")
    entries = flatten(info) if info.get("entries") is not None else [info]
    source_name = info.get("uploader") or info.get("channel") or info.get("title") or url
    vids = [{"url": entry_url(e, platform), "title": e.get("title") or e.get("description") or "",
             "views": e.get("view_count") or 0, "duration": e.get("duration")} for e in entries]
    vids = [v for v in vids if v["url"]]
    if sort == "popular":
        vids.sort(key=lambda v: v["views"] or 0, reverse=True)
    if limit:
        vids = vids[:limit]
    return source_name, vids


def clean_caption_text(raw, ext):
    if ext == "json3":
        try:
            data = json.loads(raw)
            parts = []
            for ev in data.get("events", []):
                t = "".join(s.get("utf8", "") for s in ev.get("segs") or [])
                if t.strip():
                    parts.append(t.strip())
            return re.sub(r"\s+", " ", " ".join(parts)).strip()
        except Exception:
            pass
    lines, recent = [], []
    for line in raw.splitlines():
        s = line.strip()
        if not s or "-->" in s or s.isdigit() or s.startswith(("WEBVTT", "Kind:", "Language:", "NOTE", "STYLE")):
            continue
        s = re.sub(r"<[^>]+>", "", s)
        s = re.sub(r"\{[^}]*\}", "", s).strip()
        if not s or s in recent:
            continue
        lines.append(s)
        recent = (recent + [s])[-3:]
    return re.sub(r"\s+", " ", " ".join(lines)).strip()


def pick_caption(info, lang):
    """Choisit la meilleure piste de sous-titres. Retourne (url, ext) ou None."""
    def best_fmt(tracks):
        by = {t.get("ext"): t for t in tracks if t.get("url")}
        for ext in ("json3", "vtt", "srt", "srv3", "ttml"):
            if ext in by:
                return by[ext]["url"], ext
        if tracks and tracks[0].get("url"):
            return tracks[0]["url"], tracks[0].get("ext", "vtt")
        return None

    subs = info.get("subtitles") or {}
    auto = info.get("automatic_captions") or {}
    subs = {k: v for k, v in subs.items() if k != "live_chat"}
    orig_lang = (info.get("language") or "").split("-")[0]
    order = []
    if lang:
        order += [k for k in subs if k.split("-")[0] == lang]
    if orig_lang:
        order += [k for k in subs if k.split("-")[0] == orig_lang]
    order += list(subs.keys())
    for k in order:
        r = best_fmt(subs[k])
        if r:
            return r
    # sous-titres auto : uniquement la langue d'origine (pas les traductions auto)
    keys = [k for k in auto if k.endswith("-orig")]
    if lang:
        keys += [k for k in auto if k.split("-")[0] == lang and not ("-" in k and not k.endswith("-orig"))]
    if orig_lang:
        keys += [k for k in auto if k == orig_lang]
    for k in keys:
        r = best_fmt(auto[k])
        if r:
            return r
    return None


def groq_transcribe(path, lang):
    data = {"model": GROQ_MODEL, "response_format": "text", "temperature": "0"}
    if lang:
        data["language"] = lang
    for attempt in range(6):
        with open(path, "rb") as f:
            r = requests.post("https://api.groq.com/openai/v1/audio/transcriptions",
                              headers={"Authorization": f"Bearer {GROQ_KEY}"},
                              data=data, files={"file": (os.path.basename(path), f)}, timeout=300)
        if r.status_code == 200:
            return r.text.strip()
        if r.status_code in (429, 500, 502, 503):
            wait = r.headers.get("retry-after")
            try:
                wait = float(wait)
            except (TypeError, ValueError):
                wait = 15 * (attempt + 1)
            time.sleep(min(max(wait, 3), 120))
            continue
        if r.status_code == 401:
            raise RuntimeError("Clé Groq invalide.")
        raise RuntimeError(f"Groq a répondu {r.status_code} : {r.text[:150]}")
    raise RuntimeError("Limite Groq atteinte, réessaie plus tard.")


def transcribe_audio(platform, url, lang, tmp):
    opts = ydl_opts(platform, format="bestaudio/worst[ext=mp4]/worst/best",
                    outtmpl=os.path.join(tmp, "media.%(ext)s"))
    with yt_dlp.YoutubeDL(opts) as y:
        y.extract_info(url, download=True)
    media = [p for p in glob.glob(os.path.join(tmp, "media.*")) if not p.endswith(".part")]
    if not media:
        raise RuntimeError("Téléchargement de l'audio impossible.")
    # audio mono 16 kHz très léger, découpé en morceaux de 20 min
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", media[0], "-vn", "-ac", "1", "-ar", "16000",
                    "-b:a", "32k", "-f", "segment", "-segment_time", "1200",
                    os.path.join(tmp, "part_%03d.mp3")], check=True)
    parts = sorted(glob.glob(os.path.join(tmp, "part_*.mp3")))
    if not parts:
        raise RuntimeError("Cette vidéo n'a pas de son.")
    return " ".join(t for t in (groq_transcribe(p, lang) for p in parts) if t).strip()


def is_bot_block(msg):
    m = msg.lower()
    return "sign in to confirm" in m or "not a bot" in m or "cookies" in m and "youtube" in m


def bot_message():
    if YT["status"] == "absent":
        return "YouTube bloque le serveur : ajoute tes cookies YouTube (bouton en haut)."
    YT["status"] = "expired"
    return "Tes cookies YouTube ont expiré : refais l'export et recolle-les."


def process_item(job, item):
    if job.get("blocked", 0) >= 3:  # YouTube bloque déjà : inutile d'insister
        item["status"], item["msg"] = "error", job["block_msg"]
        return
    item["status"] = "working"
    tmp = tempfile.mkdtemp(dir=BASE)
    try:
        with yt_dlp.YoutubeDL(ydl_opts(job["platform"], skip_download=True)) as y:
            info = y.extract_info(item["url"], download=False)
            item["title"] = info.get("title") or info.get("description") or item["title"]
            text = ""
            cap = pick_caption(info, job["lang"])
            if cap:
                try:
                    raw = y.urlopen(cap[0]).read().decode("utf-8", "ignore")
                    text = clean_caption_text(raw, cap[1])
                    item["method"] = "sous-titres"
                except Exception:
                    text = ""
        if len(text) < 20:
            if not GROQ_KEY:
                raise RuntimeError("Pas de sous-titres et aucune clé Groq configurée.")
            item["msg"] = "Transcription IA…"
            text = transcribe_audio(job["platform"], item["url"], job["lang"], tmp)
            item["method"] = "IA"
        if not text:
            item["status"], item["msg"] = "empty", "Aucune parole détectée"
        else:
            item["text"], item["status"], item["msg"] = text, "done", ""
        if job["platform"] == "youtube" and YT["status"] != "absent":
            YT["status"] = "ok"
    except Exception as e:
        msg = re.sub(r"\x1b\[[0-9;]*m", "", str(e)).replace("ERROR: ", "")
        if job["platform"] == "youtube" and is_bot_block(msg):
            msg = bot_message()
            job["blocked"] = job.get("blocked", 0) + 1
            job["block_msg"] = msg
        item["status"], item["msg"] = "error", msg[:200]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def paragraphs(text, n=4):
    sentences = re.split(r"(?<=[.!?…])\s+", text.strip())
    if len(sentences) < 2:  # pas de ponctuation (sous-titres auto) : couper tous les ~70 mots
        words = text.split()
        return [" ".join(words[i:i + 70]) for i in range(0, len(words), 70)]
    return [" ".join(sentences[i:i + n]) for i in range(0, len(sentences), n)]


def build_docx(job, path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    is_file = job["platform"] == "file"
    doc.add_heading(f"{'Transcription' if is_file else 'Scripts'} — {job['source_name']}", 0)
    done = [i for i in job["items"] if i["status"] == "done"]
    meta = doc.add_paragraph()
    label = {"tiktok": "TikTok", "youtube": "YouTube", "file": "Fichier importé"}.get(job["platform"], "")
    count = "transcription intégrale" if job["platform"] == "file" else f"{len(done)} scripts"
    r = meta.add_run(f"{label} · {count} · "
                     f"{datetime.datetime.now().strftime('%d/%m/%Y')}")
    r.font.color.rgb, r.font.size = RGBColor(0x66, 0x66, 0x66), Pt(10)
    for n, it in enumerate(done, 1):
        if is_file:
            for p in paragraphs(it["text"]):
                doc.add_paragraph(p)
            continue
        doc.add_heading(f"Script {n}", 1)
        info = doc.add_paragraph()
        t = info.add_run((it["title"] or "")[:200])
        t.italic = True
        info.add_run("\n" + it["url"]).font.size = Pt(9)
        for p in paragraphs(it["text"]):
            doc.add_paragraph(p)
    failed = [i for i in job["items"] if i["status"] != "done"]
    if failed:
        doc.add_page_break()
        doc.add_heading("Vidéos non transcrites", 1)
        for it in failed:
            doc.add_paragraph(f"{it['url']} — {it['msg']}", style="List Bullet")
    doc.save(path)


def run_job(job):
    try:
        job["phase"] = "listing"
        name, vids = list_videos(job["platform"], job["url"], job["limit"], job["sort"])
        if not vids:
            raise RuntimeError("Aucune vidéo trouvée sur ce compte.")
        job["source_name"] = name
        job["items"] = [{"n": i + 1, "url": v["url"], "title": v["title"], "status": "pending",
                         "msg": "", "method": "", "text": ""} for i, v in enumerate(vids)]
        job["phase"] = "transcribing"
        with ThreadPoolExecutor(max_workers=WORKERS) as ex:
            list(ex.map(lambda it: process_item(job, it), job["items"]))
        if not any(i["status"] == "done" for i in job["items"]):
            raise RuntimeError(job.get("block_msg") or "Aucune vidéo n'a pu être transcrite.")
        safe = re.sub(r"[^\w-]+", "_", name)[:40].strip("_") or "scripts"
        job["filename"] = f"Scripts_{safe}.docx"
        job["file"] = os.path.join(BASE, job["id"] + ".docx")
        build_docx(job, job["file"])
        job["phase"] = "done"
    except Exception as e:
        msg = re.sub(r"\x1b\[[0-9;]*m", "", str(e)).replace("ERROR: ", "")
        if job["platform"] == "youtube" and is_bot_block(msg):
            msg = bot_message()
        job["phase"], job["error"] = "error", msg[:300]
    job["finished"] = time.time()


def cleanup():
    now = time.time()
    with LOCK:
        for jid in list(JOBS):
            j = JOBS[jid]
            if j.get("finished") and now - j["finished"] > 6 * 3600:
                if j.get("file") and os.path.exists(j["file"]):
                    os.remove(j["file"])
                del JOBS[jid]


# ---------------------------------------------------------------- routes
def authed():
    return not APP_PASSWORD or session.get("ok")


@app.route("/login", methods=["GET", "POST"])
def login():
    err = ""
    if request.method == "POST":
        if request.form.get("password") == APP_PASSWORD:
            session["ok"] = True
            session.permanent = True
            return redirect("/")
        err = "Mot de passe incorrect"
    return render_template("login.html", err=err)


@app.route("/")
def index():
    if not authed():
        return redirect("/login")
    return render_template("index.html", groq=bool(GROQ_KEY), yt=YT["status"], yt_src=YT["source"])


@app.post("/api/start")
def start():
    if not authed():
        return jsonify(error="Non connecté"), 401
    cleanup()
    d = request.get_json(force=True)
    platform = "tiktok" if d.get("platform") == "tiktok" else "youtube"
    raw = (d.get("source") or "").strip()
    if not raw:
        return jsonify(error="Indique un compte ou un lien."), 400
    try:
        limit = int(d.get("limit") or 0)
    except ValueError:
        limit = 0
    job = {"id": uuid.uuid4().hex[:12], "platform": platform, "source": raw,
           "url": normalize_source(platform, raw, d.get("yt_tab")),
           "limit": max(0, min(limit, 1000)), "sort": "popular" if d.get("sort") == "popular" else "recent",
           "lang": (d.get("lang") or "").strip()[:5], "phase": "queued", "items": [],
           "source_name": raw, "error": "", "created": time.time()}
    with LOCK:
        JOBS[job["id"]] = job
    threading.Thread(target=run_job, args=(job,), daemon=True).start()
    return jsonify(job_id=job["id"])


@app.get("/api/job/<jid>")
def job_status(jid):
    if not authed():
        return jsonify(error="Non connecté"), 401
    j = JOBS.get(jid)
    if not j:
        return jsonify(error="Tâche introuvable (le serveur a peut-être redémarré)."), 404
    items = [{k: v for k, v in i.items() if k != "text"} | {"chars": len(i["text"])} for i in j["items"]]
    return jsonify(id=j["id"], phase=j["phase"], error=j["error"], source_name=j["source_name"],
                   platform=j["platform"], items=items, has_file=bool(j.get("file")),
                   part=j.get("part", 0), parts=j.get("parts", 0), duration=j.get("duration", 0))


@app.get("/api/job/<jid>/text")
def job_text(jid):
    if not authed():
        return "Non connecté", 401
    j = JOBS.get(jid)
    if not j:
        return "Introuvable", 404
    done = [i for i in j["items"] if i["status"] == "done"]
    out = "\n\n".join(f"Script {n}\n{i['text']}" for n, i in enumerate(done, 1))
    return out, 200, {"Content-Type": "text/plain; charset=utf-8"}


@app.get("/api/job/<jid>/docx")
def job_docx(jid):
    if not authed():
        return "Non connecté", 401
    j = JOBS.get(jid)
    if not j or not j.get("file") or not os.path.exists(j["file"]):
        return "Fichier introuvable", 404
    return send_file(j["file"], as_attachment=True, download_name=j["filename"])


# ------------------------------------------------ transcription d'un fichier
def probe_duration(path):
    try:
        out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration",
                              "-of", "default=nw=1:nk=1", path], capture_output=True, text=True, timeout=60)
        return float(out.stdout.strip())
    except Exception:
        return 0.0


def run_file_job(job, src):
    item = job["items"][0]
    tmp = os.path.dirname(src)
    try:
        job["phase"] = "transcribing"
        item["status"], item["msg"] = "working", "Extraction du son…"
        job["duration"] = probe_duration(src)
        # son mono 16 kHz léger, en morceaux de 10 min (bien sous la limite Groq)
        r = subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", src, "-vn", "-ac", "1", "-ar", "16000",
                            "-b:a", "32k", "-f", "segment", "-segment_time", "600",
                            os.path.join(tmp, "part_%03d.mp3")], capture_output=True, text=True)
        try:
            os.remove(src)  # libère le disque tout de suite
        except OSError:
            pass
        parts = sorted(glob.glob(os.path.join(tmp, "part_*.mp3")))
        if r.returncode != 0 or not parts:
            raise RuntimeError("Impossible de lire le son de ce fichier. Vérifie que c'est bien un audio ou une vidéo.")
        job["parts"] = len(parts)
        texts = []
        for i, p in enumerate(parts, 1):
            item["msg"] = f"Transcription IA · partie {i}/{len(parts)}"
            job["part"] = i - 1
            t = groq_transcribe(p, job["lang"])
            if t:
                texts.append(t)
            job["part"] = i
        text = " ".join(texts).strip()
        if not text:
            item["status"], item["msg"] = "empty", "Aucune parole détectée dans ce fichier."
            raise RuntimeError("Aucune parole détectée dans ce fichier.")
        item.update(text=text, status="done", msg="", method="IA")
        safe = re.sub(r"[^\w-]+", "_", os.path.splitext(job["source_name"])[0])[:40].strip("_") or "transcription"
        job["filename"] = f"Transcription_{safe}.docx"
        job["file"] = os.path.join(BASE, job["id"] + ".docx")
        build_docx(job, job["file"])
        job["phase"] = "done"
    except Exception as e:
        msg = re.sub(r"\x1b\[[0-9;]*m", "", str(e))
        if item["status"] == "working":
            item["status"], item["msg"] = "error", msg[:200]
        job["phase"], job["error"] = "error", msg[:300]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
        job["finished"] = time.time()


@app.post("/api/file")
def upload_file():
    if not authed():
        return jsonify(error="Non connecté"), 401
    if not GROQ_KEY:
        return jsonify(error="La clé Groq n'est pas configurée : la transcription IA est impossible."), 400
    cleanup()
    f = request.files.get("file")
    if not f or not f.filename:
        return jsonify(error="Choisis un fichier audio ou vidéo."), 400
    jid = uuid.uuid4().hex[:12]
    tmp = tempfile.mkdtemp(dir=BASE)
    ext = os.path.splitext(f.filename)[1][:10] or ".bin"
    src = os.path.join(tmp, "input" + ext)
    f.save(src)
    name = os.path.basename(f.filename)[:120]
    job = {"id": jid, "platform": "file", "source": name, "url": "", "limit": 1, "sort": "",
           "lang": (request.form.get("lang") or "").strip()[:5], "phase": "queued", "error": "",
           "source_name": name, "created": time.time(), "part": 0, "parts": 0, "duration": 0,
           "items": [{"n": 1, "url": "", "title": name, "status": "pending", "msg": "En attente",
                      "method": "", "text": ""}]}
    with LOCK:
        JOBS[jid] = job
    threading.Thread(target=run_file_job, args=(job, src), daemon=True).start()
    return jsonify(job_id=jid)


@app.get("/api/job/<jid>/full")
def job_full(jid):
    if not authed():
        return jsonify(error="Non connecté"), 401
    j = JOBS.get(jid)
    if not j:
        return jsonify(error="Tâche introuvable"), 404
    it = j["items"][0] if j["items"] else {}
    return jsonify(text=it.get("text", ""))


@app.get("/api/cookies")
def cookies_status():
    if not authed():
        return jsonify(error="Non connecté"), 401
    return jsonify(status=YT["status"], source=YT["source"], updated=YT["updated"])


@app.post("/api/cookies")
def cookies_upload():
    if not authed():
        return jsonify(error="Non connecté"), 401
    f = request.files.get("file")
    text = f.read().decode("utf-8", "ignore") if f else (request.form.get("text") or "")
    lines = [l for l in text.splitlines() if l.strip() and not l.startswith("#")]
    domains = sorted({l.split("\t")[0] for l in lines if "\t" in l})[:8]
    print(f"Dépôt cookies : {len(text)} caractères, {len(lines)} lignes, domaines {domains}, "
          f"1re ligne : {text[:40]!r}", flush=True)
    if text.lstrip().startswith(("[", "{")):
        return jsonify(error="Ce fichier est au format JSON. Dans l'extension, choisis le format « Netscape » puis refais Export."), 400
    if lines and not any("youtube.com" in l for l in lines):
        return jsonify(error="Ce fichier ne contient pas de cookies YouTube : fais l'export en étant sur youtube.com."), 400
    if not valid_yt_cookies(text):
        return jsonify(error="Ce fichier ne contient pas de cookies YouTube. Exporte-le depuis youtube.com avec l'extension."), 400
    install_yt_cookies(text, "upload")
    return jsonify(status="ok")


@app.get("/favicon.ico")
def favicon():
    return send_file(os.path.join(app.root_path, "static", "favicon.ico"), mimetype="image/x-icon")


@app.errorhandler(413)
def too_big(_):
    return jsonify(error="Fichier trop lourd (2 Go maximum)."), 413


@app.get("/health")
def health():
    return "ok"


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=True)
